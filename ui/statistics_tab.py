"""Statistics tab.

Receives per-trial metric rows from the Analyze tab (via the shared
``ResultsStore``) and/or from an imported Excel export, and shows them as a
table. Statistical tests are a future addition; for now this realises the
Analyze -> Statistics hand-off and lets the user see/collect results across
trials.

It deliberately omits ``supports_load_file`` / project hooks — File ▸ project
actions don't apply here; the Statistics menu provides Import/Clear.
"""

from PyQt6.QtWidgets import (
    QWidget, QHBoxLayout, QVBoxLayout, QFrame, QLabel, QTableWidget,
    QTableWidgetItem, QStackedLayout, QFileDialog, QMessageBox,
)
from PyQt6.QtCore import Qt

from ui import style as S


class StatisticsTab(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.results_store = None

        root = QHBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        # --- sidebar stub (mirrors the Analyze tab's chrome) ---
        self._sidebar = QFrame()
        self._sidebar.setObjectName("sidebar")
        self._sidebar.setFixedWidth(220)
        sb = QVBoxLayout(self._sidebar)
        sb.setContentsMargins(14, 14, 14, 14)
        sb.setSpacing(8)
        self._side_title = QLabel("Statistics")
        self._side_title.setObjectName("section-toggle")
        sb.addWidget(self._side_title)
        self._side_note = QLabel(
            "Results from Analyze land here.\n"
            "Menu ▸ Statistics:\n"
            "• Import results (Excel)…\n"
            "• Clear results\n\n"
            "Statistical tests: coming soon."
        )
        self._side_note.setObjectName("field-label")
        self._side_note.setWordWrap(True)
        sb.addWidget(self._side_note)
        self._count_label = QLabel("0 rows")
        self._count_label.setObjectName("field-label")
        sb.addWidget(self._count_label)
        sb.addStretch(1)

        # --- content: empty-placeholder OR table ---
        self._content = QWidget()
        self._stack = QStackedLayout(self._content)

        self._empty = QWidget()
        el = QVBoxLayout(self._empty)
        el.addStretch(1)
        self._headline = QLabel("Statistics")
        self._headline.setAlignment(Qt.AlignmentFlag.AlignCenter)
        el.addWidget(self._headline)
        self._sub = QLabel(
            "아직 결과가 없습니다.\n\n"
            "Analyze 탭에서 분석 후 메뉴 ▸ Analyze ▸ \"Send results → Statistics\"\n"
            "또는 메뉴 ▸ Statistics ▸ \"Import results (Excel)…\" 로 불러오세요."
        )
        self._sub.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._sub.setWordWrap(True)
        el.addWidget(self._sub)
        el.addStretch(2)

        self._table = QTableWidget()
        self._table.setObjectName("results-table")
        self._table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self._table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self._table.setAlternatingRowColors(True)

        self._stack.addWidget(self._empty)
        self._stack.addWidget(self._table)

        root.addWidget(self._sidebar)
        root.addWidget(self._content, 1)

        self.apply_theme()

    # ----- shared store wiring -----
    def set_results_store(self, store):
        self.results_store = store
        if store is not None:
            store.add_listener(self._refresh_table)
        self._refresh_table()

    def _refresh_table(self):
        store = self.results_store
        if store is None or store.is_empty():
            self._count_label.setText("0 rows")
            self._stack.setCurrentWidget(self._empty)
            return
        headers, rows = store.table()
        self._table.clear()
        self._table.setColumnCount(len(headers))
        self._table.setHorizontalHeaderLabels(headers)
        self._table.setRowCount(len(rows))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                if isinstance(val, float):
                    text = f"{val:.3f}"
                else:
                    text = "" if val is None else str(val)
                self._table.setItem(r, c, QTableWidgetItem(text))
        self._table.resizeColumnsToContents()
        self._count_label.setText(f"{len(rows)} rows")
        self._stack.setCurrentWidget(self._table)

    # ----- menu actions (called by MainWindow) -----
    def import_results_excel(self):
        if self.results_store is None:
            return
        path, _ = QFileDialog.getOpenFileName(
            self, "Import results", "", "Excel Files (*.xlsx)"
        )
        if not path:
            return
        try:
            rows = self._read_results_sheet(path)
        except ModuleNotFoundError:
            QMessageBox.critical(self, "Missing package", "openpyxl is required to read Excel files.")
            return
        except Exception as exc:
            QMessageBox.warning(self, "Import failed", f"Could not read results:\n{exc}")
            return
        if not rows:
            QMessageBox.information(self, "No rows", "No 'Results' sheet rows found.")
            return
        self.results_store.add_rows(rows, source=path)

    def clear_results(self):
        if self.results_store is not None:
            self.results_store.clear()

    @staticmethod
    def _read_results_sheet(path):
        """Read the 'Results' sheet of an exported workbook into store rows."""
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Results"] if "Results" in wb.sheetnames else wb.active
        it = ws.iter_rows(values_only=True)
        try:
            header = list(next(it))
        except StopIteration:
            return []
        header = [str(h) if h is not None else "" for h in header]
        rows = []
        for raw in it:
            if raw is None or all(v is None for v in raw):
                continue
            row = {}
            for col, val in zip(header, raw):
                if not col:
                    continue
                row[col] = val
            if row:
                rows.append(row)
        wb.close()
        return rows

    def apply_theme(self):
        self._content.setStyleSheet(f"background: {S.BG_DARK};")
        self._headline.setStyleSheet(
            f"color: {S.TEXT_PRIMARY}; font-size: 24px; font-weight: 700;"
        )
        self._sub.setStyleSheet(f"color: {S.TEXT_SECONDARY}; font-size: 13px;")
